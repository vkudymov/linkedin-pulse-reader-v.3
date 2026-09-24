from __future__ import annotations

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from ..types import ReportItem, ReportMeta


def _autosize(ws) -> None:  # type: ignore[no-untyped-def]
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        max_len = 0
        for cell in ws[letter]:
            if cell.value is None:
                continue
            max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max(12, max_len + 2), 60)


def build_xlsx(*, items: list[ReportItem], meta: ReportMeta) -> bytes:
    matched = [i for i in items if i.analyzed and i.match is True]
    not_matched = [i for i in items if i.analyzed and i.match is False]
    not_analyzed = [i for i in items if not i.analyzed]

    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Summary"
    ws0.append(["kind", meta.kind])
    ws0.append(["search_title", meta.search_title])
    ws0.append(["run_at_iso", meta.run_at_iso or ""])
    ws0.append(["min_score", meta.min_score if meta.min_score is not None else ""])
    ws0.append(["target_found", meta.target_found if meta.target_found is not None else ""])
    ws0.append(["fetched_count", meta.fetched_count if meta.fetched_count is not None else ""])
    ws0.append(["analyzed_count", meta.analyzed_count if meta.analyzed_count is not None else ""])
    ws0.append(["matched_count", meta.matched_count if meta.matched_count is not None else ""])
    _autosize(ws0)

    def add_sheet(name: str, rows: list[ReportItem]) -> None:
        ws = wb.create_sheet(name)
        ws.append(
            [
                "title",
                "url",
                "company",
                "location",
                "score",
                "match",
                "reason",
                "matched_requirements",
                "missing_requirements",
                "red_flags",
                "comment_text",
                "text",
                "description",
            ]
        )
        for i in rows:
            ws.append(
                [
                    i.title or "",
                    i.url or "",
                    i.company or "",
                    i.location or "",
                    i.score if i.score is not None else "",
                    i.match if i.match is not None else "",
                    i.reason or "",
                    "; ".join(list(i.matched_requirements or ())),
                    "; ".join(list(i.missing_requirements or ())),
                    "; ".join(list(i.red_flags or ())),
                    i.comment_text or "",
                    i.text or "",
                    i.description or "",
                ]
            )
        _autosize(ws)

    add_sheet("Matched", matched)
    add_sheet("NotMatched", not_matched)
    add_sheet("NotAnalyzed", not_analyzed)

    import io

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

