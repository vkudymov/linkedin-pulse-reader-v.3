from __future__ import annotations

import io

from docx import Document

from ..types import ReportItem, ReportMeta


def build_docx(*, items: list[ReportItem], meta: ReportMeta) -> bytes:
    matched = [i for i in items if i.analyzed and i.match is True]
    not_matched = [i for i in items if i.analyzed and i.match is False]
    not_analyzed = [i for i in items if not i.analyzed]

    doc = Document()
    doc.add_heading(f"Search report ({meta.kind})", level=1)
    doc.add_paragraph(f"Title: {meta.search_title}")
    if meta.run_at_iso:
        doc.add_paragraph(f"Run at: {meta.run_at_iso}")
    if meta.min_score is not None:
        doc.add_paragraph(f"Min score: {meta.min_score}")
    if meta.target_found is not None:
        doc.add_paragraph(f"Target found: {meta.target_found}")
    if meta.fetched_count is not None:
        doc.add_paragraph(f"Fetched: {meta.fetched_count}")
    if meta.analyzed_count is not None:
        doc.add_paragraph(f"Analyzed: {meta.analyzed_count}")
    if meta.matched_count is not None:
        doc.add_paragraph(f"Matched: {meta.matched_count}")

    def section(title: str, rows: list[ReportItem]) -> None:
        doc.add_heading(title, level=2)
        if not rows:
            doc.add_paragraph("(empty)")
            return
        for idx, i in enumerate(rows, start=1):
            doc.add_heading(f"{idx}. {i.title or '<no title>'}", level=3)
            if i.url:
                doc.add_paragraph(f"URL: {i.url}")
            if i.company:
                doc.add_paragraph(f"Company: {i.company}")
            if i.location:
                doc.add_paragraph(f"Location: {i.location}")
            if i.score is not None:
                doc.add_paragraph(f"Score: {i.score}")
            if i.reason:
                doc.add_paragraph(f"Reason: {i.reason}")
            if i.matched_requirements:
                doc.add_paragraph("Matched requirements: " + "; ".join(list(i.matched_requirements)))
            if i.missing_requirements:
                doc.add_paragraph("Missing requirements: " + "; ".join(list(i.missing_requirements)))
            if i.red_flags:
                doc.add_paragraph("Red flags: " + "; ".join(list(i.red_flags)))
            if i.comment_text:
                doc.add_paragraph("Comment:")
                doc.add_paragraph(i.comment_text)
            if i.text:
                doc.add_paragraph("Text:")
                doc.add_paragraph(i.text)
            elif i.description:
                doc.add_paragraph("Description:")
                doc.add_paragraph(i.description)

    section("MATCHED", matched)
    section("NOT_MATCHED", not_matched)
    section("NOT_ANALYZED", not_analyzed)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

